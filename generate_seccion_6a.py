# -*- coding: utf-8 -*-
"""
Script to generate and append Section 6 and 6.a (Análisis de Datos)
to Finanzas Trabajo Parcial.docx with exact consistent formatting.
"""
import os
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_seccion_6a(doc_path, output_path):
    doc = docx.Document(doc_path)
    
    # Colors
    UPC_RED = RGBColor(200, 16, 46)     # #c8102e
    DARK_GRAY = RGBColor(67, 67, 67)    # #434343
    WHITE = RGBColor(255, 255, 255)
    
    # Helper to add page break if needed, or add spacing
    # Let's add a clean paragraph separator
    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(12)
    p_sep.paragraph_format.space_after = Pt(6)
    
    # -------------------------------------------------------------
    # 1. Main Heading 1: 6. ANÁLISIS Y DISEÑO DEL SISTEMA
    # -------------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.style = 'Heading 1'
    h1.paragraph_format.space_before = Pt(21)
    h1.paragraph_format.space_after = Pt(9)
    run_h1 = h1.add_run('6. ANÁLISIS Y DISEÑO DEL SISTEMA')
    run_h1.font.size = Pt(22)
    run_h1.font.bold = True
    run_h1.font.color.rgb = UPC_RED
    
    p_intro6 = doc.add_paragraph()
    p_intro6.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro6.paragraph_format.space_after = Pt(8)
    p_intro6.paragraph_format.line_spacing = 1.15
    p_intro6.add_run(
        "En este capítulo se desarrolla la especificación técnica, el modelo de datos y los fundamentos "
        "matemático-financieros que sustentan la aplicación web para el control permanente de cuenta corriente "
        "y otorgamiento de créditos comerciales de barrio. El sistema ha sido concebido bajo un modelo multiempresa "
        "y multirol, posibilitando que diversos comercios de una misma zona de influencia (como bodegas, panaderías, "
        "carnicerías, pescaderías, pollerías, bazares, boticas, licorerías, peluquerías o spas) gestionen de forma "
        "autónoma e independiente sus clientes, catálogos y políticas crediticias sobre una plataforma tecnológica compartida.\n\n"
        "Asimismo, el diseño computacional se adhiere con estricta rigurosidad a las convenciones de cálculo del curso "
        "(base anual de 360 días, mes comercial de 30 días, amortización bajo el método francés vencido simple ordinario, "
        "capitalización íntegra de intereses en períodos de gracia total, y orden legal de imputación o prelación de pagos: "
        "mora, compensatorio y capital). Todo el análisis se formula desde la perspectiva del comercio que otorga el crédito, "
        "garantizando que la herramienta proporcione información certera sobre las cuentas por cobrar, mitigue el riesgo de "
        "incumplimiento mediante límites de endeudamiento y ofrezca indicadores financieros avanzados (TIR, VAN, TCEA) para "
        "evaluar la rentabilidad del financiamiento otorgado frente al costo de oportunidad del negocio."
    )
    
    # -------------------------------------------------------------
    # 2. Heading 2: a) Análisis de Datos
    # -------------------------------------------------------------
    h2 = doc.add_paragraph()
    h2.style = 'Heading 2'
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    run_h2 = h2.add_run('a) Análisis de Datos')
    run_h2.font.size = Pt(16)
    run_h2.font.bold = True
    
    p_intro_a = doc.add_paragraph()
    p_intro_a.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro_a.paragraph_format.space_after = Pt(8)
    p_intro_a.paragraph_format.line_spacing = 1.15
    p_intro_a.add_run(
        "El análisis de datos permite estructurar formalmente el flujo de información que transita a través de la aplicación. "
        "Para responder fielmente a los requerimientos de la ingeniería de software y a la rúbrica de evaluación del curso, "
        "las variables del sistema se han dividido en tres categorías exhaustivas:\n"
        "1. Datos de Entrada: Parámetros del sistema, datos de configuración de tiendas, perfiles de acceso, información "
        "del catálogo de bienes y servicios, condiciones contractuales pactadas con cada cliente y eventos transaccionales de venta.\n"
        "2. Datos de Salida: Reportes y liquidaciones producidos para el comerciante y el cliente, incluyendo el listado de pago "
        "mensual a la fecha de corte, el cronograma detallado de cuotas francesas, los estados de cuenta corriente y las métricas "
        "financieras de costo y rendimiento (TCEA, TIR, VAN).\n"
        "3. Datos Intermedios: Variables auxiliares, factores de conversión de tasas, cómputo de días calendario y de gracia, "
        "deuda capitalizada, componentes de amortización y vectores de flujos netos generados por los motores de cálculo del sistema."
    )
    
    # Table styling helpers
    def format_cell(cell, fill_hex, border_color="bfbfbf", border_sz="4"):
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>\n'
            f'  <w:left w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>\n'
            f'  <w:bottom w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>\n'
            f'  <w:right w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)
        if fill_hex:
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}" w:val="clear"/>')
            tcPr.append(shd)
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>\n'
            f'  <w:top w:w="80" w:type="dxa"/>\n'
            f'  <w:left w:w="110" w:type="dxa"/>\n'
            f'  <w:bottom w:w="80" w:type="dxa"/>\n'
            f'  <w:right w:w="110" w:type="dxa"/>\n'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)
        vAlign = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="top"/>')
        tcPr.append(vAlign)

    def set_table_layout(table, col_widths, total_dxa=9360):
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="nil" w:sz="0" w:space="0" w:color="000000"/>\n'
            f'  <w:left w:val="nil" w:sz="0" w:space="0" w:color="000000"/>\n'
            f'  <w:bottom w:val="nil" w:sz="0" w:space="0" w:color="000000"/>\n'
            f'  <w:right w:val="nil" w:sz="0" w:space="0" w:color="000000"/>\n'
            f'  <w:insideH w:val="nil" w:sz="0" w:space="0" w:color="000000"/>\n'
            f'  <w:insideV w:val="nil" w:sz="0" w:space="0" w:color="000000"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)
        tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:w="{total_dxa}" w:type="dxa"/>')
        tblPr.append(tblW)
        layout = parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>')
        tblPr.append(layout)
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = w

    # -------------------------------------------------------------
    # 3. Datos de Entrada
    # -------------------------------------------------------------
    sub1 = doc.add_paragraph()
    sub1.paragraph_format.space_before = Pt(12)
    sub1.paragraph_format.space_after = Pt(4)
    run_sub1 = sub1.add_run('6.a.1. Datos de Entrada')
    run_sub1.font.size = Pt(13)
    run_sub1.font.bold = True
    run_sub1.font.color.rgb = UPC_RED
    
    p_desc_in = doc.add_paragraph()
    p_desc_in.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_desc_in.paragraph_format.space_after = Pt(8)
    p_desc_in.paragraph_format.line_spacing = 1.15
    p_desc_in.add_run(
        "Los datos de entrada agrupan todas las constantes institucionales y variables operativas que ingresan al sistema. "
        "Se dividen funcionalmente en: constantes globales financieras (Anexo A), datos de registro del comercio, credenciales "
        "de acceso multirol, catálogo de bienes/servicios, condiciones de la línea de crédito del cliente, registro de compras "
        "a crédito en el punto de venta y eventos de liquidación de pagos. A continuación se presenta la tabla detallada:"
    )
    
    # Columns: VARIABLE / CONSTANTE | DESCRIPCIÓN | TIPO | RANGO / TAMAÑO | FORMATO | VALOR DEFECTO | RESTRICCIONES
    in_col_widths = [1600, 2260, 700, 1100, 1100, 1100, 1500]  # total = 9360
    
    data_inputs = [
        # (Variable, Descripcion, Tipo, Rango/Tam, Formato, Defecto, Restricciones)
        # Constantes
        ("BASE_DIAS_ANUAL", "Base anual de cálculo comercial para tasas e intereses", "Constante Int", "360", "#", "360", "Inmutable = 360 días (Anexo A)"),
        ("DIAS_MES_COMERCIAL", "Duración fija del mes comercial para cronogramas", "Constante Int", "30", "#", "30", "Inmutable = 30 días (Anexo A)"),
        ("DECIMALES_MONEDA", "Precisión decimal obligatoria en importes dinerarios", "Constante Int", "2", "#", "2", "Redondeo a 2 decimales (Anexo A)"),
        ("DECIMALES_TASAS", "Precisión mínima para cálculo de tasas de interés", "Constante Int", "7", "#", "7", "Mínimo 7 decimales (Anexo A)"),
        ("ORDEN_PRELACION_PAGOS", "Jerarquía legal de aplicación de abonos", "Constante Array", "3 ítems", "[Mora, Comp., Cap.]", "[1, 2, 3]", "Inmutable según normativa"),
        # Tienda
        ("id_tienda", "Identificador único del comercio en la plataforma", "Int / UUID", "> 0", "#", "Autogenerado", "Clave primaria, único"),
        ("razon_social", "Razón social o nombre legal de la tienda", "String", "1 a 100 caracteres", "Texto", "Vacío", "No nulo, obligatorio"),
        ("nombre_comercial", "Nombre comercial con que el cliente identifica el negocio", "String", "1 a 80 caracteres", "Texto", "Vacío", "No nulo, obligatorio"),
        ("ruc_tienda", "Número de RUC del comercio de barrio", "String", "11 dígitos", "###########", "Vacío", "11 dígitos, inicia con 10 o 20"),
        ("giro_negocio", "Rubro comercial del establecimiento", "Enum", "{Bodega, Panadería, Carnicería, Pollería, Pescadería, Bazar, Botica, Peluquería, Spa, Licorería, Otros}", "Texto", "Bodega", "Debe pertenecer a lista"),
        ("direccion_tienda", "Ubicación física en la zona de influencia comunal", "String", "5 a 150 caracteres", "Texto", "Vacío", "Obligatorio"),
        ("telefono_tienda", "Teléfono o celular del comercio para contacto", "String", "9 dígitos", "###-###-###", "Vacío", "Numérico, obligatorio"),
        ("correo_tienda", "Correo electrónico oficial de la tienda", "String", "Hasta 80 caracteres", "email@dominio.com", "Vacío", "Formato de correo válido"),
        ("estado_tienda", "Estado operativo del comercio en la plataforma", "Enum", "{Activa, Inactiva}", "Texto", "Activa", "Solo 'Activa' opera"),
        # Usuarios / Autenticacion
        ("usuario_login", "Nombre de usuario único para acceso al sistema", "String", "4 a 20 caracteres", "Alfanumérico", "Vacío", "Único por usuario, no nulo"),
        ("clave_acceso", "Contraseña protegida de acceso a la cuenta", "String", "8 a 32 caracteres", "Hash seguro", "Vacío", "Cifrado en base de datos"),
        ("perfil_usuario", "Rol de acceso que define permisos y vistas", "Enum", "{Admin_Sistema, Admin_Tienda, Cliente}", "Texto", "Cliente", "Obligatorio"),
        ("estado_usuario", "Estado de la cuenta de usuario", "Enum", "{Activo, Suspendido, Bloqueado}", "Texto", "Activo", "Solo 'Activo' inicia sesión"),
        # Productos / Servicios
        ("id_producto", "Código único del producto o servicio en la tienda", "String", "1 a 20 caracteres", "Alfanumérico", "Autogenerado", "Único por tienda"),
        ("descripcion_producto", "Denominación detallada del bien o servicio", "String", "3 a 100 caracteres", "Texto", "Vacío", "No nulo, obligatorio"),
        ("marca_producto", "Marca o fabricante del producto comercializado", "String", "Hasta 50 caracteres", "Texto", "Genérico", "Opcional en servicios"),
        ("proveedor_producto", "Distribuidor o proveedor del bien", "String", "Hasta 80 caracteres", "Texto", "Elaboración propia", "Opcional según rubro"),
        ("unidad_medida", "Unidad comercial de despacho del ítem", "Enum", "{Unidad, Kg, Litro, Paquete, Docena, Porción, Servicio, Otro}", "Texto", "Unidad", "Obligatorio"),
        ("precio_contado (PC)", "Precio al contado si se cancela en el acto", "Float", "0.10 <= PC <= 20,000.00", "S/ #,###.##", "0.00", "PC > 0"),
        ("precio_lista (PL)", "Precio de venta fijado para crédito comercial", "Float", "PC <= PL <= 25,000.00", "S/ #,###.##", "Calculado = PC", "PL >= PC, > 0"),
        ("modalidad_venta_permitida", "Modalidades de pago habilitadas para el producto", "Enum", "{Solo_Fin_De_Mes, Solo_Cuotas, Ambas}", "Texto", "Ambas", "Obligatorio"),
        ("imagen_producto", "Ruta de archivo o URL de la fotografía del producto", "String", "Hasta 255 caracteres", "URL / Archivo", "placeholder.png", "Opcional"),
        # Cliente y Parametros de Credito
        ("id_cliente", "Código de identificación del cliente en la tienda", "Int / UUID", "> 0", "#", "Autogenerado", "Clave primaria por tienda"),
        ("tipo_documento", "Tipo de documento oficial de identidad", "Enum", "{DNI, CE, Pasaporte}", "Texto", "DNI", "Obligatorio"),
        ("num_documento", "Número de documento de identidad del cliente", "String", "8 dígitos (DNI) / 12 (CE)", "Alfanumérico", "Vacío", "Único por tienda, obligatorio"),
        ("nombres_cliente", "Nombres de pila del cliente vecino", "String", "2 a 60 caracteres", "Texto", "Vacío", "No nulo, alfabético"),
        ("apellidos_cliente", "Apellidos paterno y materno del cliente", "String", "2 a 60 caracteres", "Texto", "Vacío", "No nulo, alfabético"),
        ("direccion_cliente", "Domicilio del cliente en la zona de influencia", "String", "5 a 150 caracteres", "Texto", "Vacío", "Obligatorio"),
        ("telefono_cliente", "Número celular o WhatsApp para notificaciones", "String", "9 dígitos", "9########", "Vacío", "Numérico de 9 dígitos"),
        ("correo_cliente", "Correo electrónico para estados de cuenta", "String", "Hasta 80 caracteres", "email@dominio.com", "Vacío", "Opcional"),
        ("moneda_credito", "Moneda pactada para la cuenta corriente", "Enum", "{Soles (S/), Dólares ($)}", "Texto", "Soles (S/)", "Obligatorio"),
        ("limite_credito (LC)", "Límite máximo de crédito autorizado al cliente", "Float", "20.00 <= LC <= 20,000.00", "S/ #,###.##", "500.00", "LC > 0; compras no deben excederlo"),
        ("plazo_maximo_meses", "Número máximo de meses para compras financiadas", "Int", "1 <= Plazo <= 24", "#", "6", "Entero positivo > 0"),
        ("dia_corte", "Día del mes en que cierra el ciclo de compras", "Int", "1 <= dia <= 28", "#", "20", "Entero entre 1 y 28"),
        ("dia_pago", "Día del mes pactado para pago obligatorio", "Int", "1 <= dia <= 28", "#", "26", "dia_pago > dia_corte en ciclo"),
        ("tipo_tasa_compensatoria", "Tipo de tasa de interés compensatoria pactada", "Enum", "{TEA, TNA}", "Texto", "TEA", "Obligatorio"),
        ("valor_tasa_compensatoria", "Valor anual de la tasa compensatoria pactada", "Float", "0.0000000% <= Tasa <= 150%", "#.#######%", "24.0000000%", "Tasa >= 0; mín. 7 decimales"),
        ("capitalizacion_compensatoria", "Frecuencia de capitalización si tasa es TNA", "Enum / Int", "{Diaria (360), Mensual (12)}", "#", "360", "Obligatorio solo si tipo = TNA"),
        ("tipo_tasa_moratoria", "Tipo de tasa de interés moratoria por atraso", "Enum", "{TEA_mora, TNA_mora}", "Texto", "TEA_mora", "Obligatorio"),
        ("valor_tasa_moratoria", "Valor anual de la tasa moratoria pactada", "Float", "0.0000000% <= Tasa <= 60%", "#.#######%", "5.0000000%", "Tasa >= 0; mín. 7 decimales"),
        ("capitalizacion_moratoria", "Frecuencia de capitalización si moratoria es TNA", "Enum / Int", "{Diaria (360), Mensual (12)}", "#", "360", "Obligatorio solo si moratoria = TNA"),
        ("tasa_descuento_cok", "Costo de oportunidad del capital del comercio", "Float", "0.0000000% <= COK <= 100%", "#.#######%", "15.0000000%", "Para cálculo de VAN y TIR (TEA)"),
        # Compra a Credito
        ("num_ticket", "Número de comprobante de venta a crédito", "String", "10 caracteres", "TKT-######", "Autogenerado", "Único correlativo por tienda"),
        ("fecha_compra", "Fecha calendario en que se efectúa la compra", "Date", "DD/MM/AAAA", "DD/MM/AAAA", "Fecha actual", "<= Fecha actual (no futura)"),
        ("hora_compra", "Hora en que se registra la compra a crédito", "Time", "00:00:00 a 23:59:59", "HH:mm:ss", "Hora actual", "Define si es pre o post corte"),
        ("modalidad_pago_compra", "Modalidad elegida por el cliente para la compra", "Enum", "{Fin_De_Mes, Cuotas_Mensuales}", "Texto", "Fin_De_Mes", "Compatible con producto"),
        ("num_cuotas_pactadas (N)", "Número de cuotas mensuales seleccionadas", "Int", "1 <= N <= plazo_maximo", "#", "1 ó 3", "Entero, 1 <= N <= plazo_max"),
        ("cantidad_comprada (Q)", "Cantidad de unidades del ítem adquiridas", "Float", "0.01 <= Q <= 1,000.00", "#,###.##", "1.00", "Q > 0"),
        # Pago Efectivo
        ("fecha_pago_efectiva", "Fecha en que el cliente realiza el abono", "Date", "DD/MM/AAAA", "DD/MM/AAAA", "Fecha actual", ">= fecha_compra; evalúa mora"),
        ("monto_abonado", "Monto de dinero pagado por el cliente", "Float", "> 0", "S/ #,###.##", "Total exigible", "No admite pagos parciales")
    ]
    
    table_in = doc.add_table(rows=len(data_inputs)+1, cols=7)
    set_table_layout(table_in, in_col_widths)
    
    headers_in = ["VARIABLE / CONSTANTE", "DESCRIPCIÓN", "TIPO", "RANGO / TAMAÑO", "FORMATO", "VALOR POR DEFECTO", "RESTRICCIONES"]
    hdr_row = table_in.rows[0]
    hdr_trPr = hdr_row._tr.get_or_add_trPr()
    hdr_trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for i, h_text in enumerate(headers_in):
        cell = hdr_row.cells[i]
        format_cell(cell, fill_hex="c8102e")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(8.0)
        r.font.color.rgb = WHITE
        
    for r_idx, row_data in enumerate(data_inputs):
        row = table_in.rows[r_idx+1]
        fill = "f2f2f2" if r_idx % 2 == 1 else "ffffff"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            format_cell(cell, fill_hex=fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.size = Pt(8.0)
            if c_idx == 0:
                r.font.bold = True
                
    # -------------------------------------------------------------
    # 4. Datos de Salida
    # -------------------------------------------------------------
    sub2 = doc.add_paragraph()
    sub2.paragraph_format.space_before = Pt(14)
    sub2.paragraph_format.space_after = Pt(4)
    run_sub2 = sub2.add_run('6.a.2. Datos de Salida')
    run_sub2.font.size = Pt(13)
    run_sub2.font.bold = True
    run_sub2.font.color.rgb = UPC_RED
    
    p_desc_out = doc.add_paragraph()
    p_desc_out.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_desc_out.paragraph_format.space_after = Pt(8)
    p_desc_out.paragraph_format.line_spacing = 1.15
    p_desc_out.add_run(
        "Los datos de salida constituyen la información calculada y presentada al usuario. Se organizan en cuatro módulos: "
        "1) Listado de Pago Mensual generado en la fecha de corte con el detalle de compras, intereses devengados e intereses por mora; "
        "2) Plan de Pagos o Cronograma de Amortización bajo el Método Francés vencido simple ordinario para compras en cuotas; "
        "3) Indicadores Financieros de Evaluación (TIR, VAN, TCEA, TEM, TED) desde la óptica del comercio; y "
        "4) Resumen de Cuenta Corriente y Línea de Crédito del cliente. A continuación se presenta la tabla formal:"
    )
    
    # Columns: VARIABLE | DESCRIPCIÓN | TIPO | CONDICIÓN / RANGO | FORMATO
    out_col_widths = [1900, 3160, 900, 1900, 1500]  # total = 9360
    
    data_outputs = [
        # Listado de pago mensual
        ("listado_compras_detalle", "Relación de compras del mes ordenadas por fecha y hora, detallando productos e importes", "Tabla", "N filas ordenadas cronológicamente", "Reporte tabular"),
        ("interes_compensatorio_acumulado", "Intereses compensatorios devengados entre fechas de compra y fecha de pago pactada", "Float", ">= 0", "S/ #,###.##"),
        ("intereses_por_mora", "Ítem detallado e independiente por mora devengada tras vencer la fecha de pago", "Float", ">= 0 (0 si canceló puntual)", "S/ #,###.##"),
        ("subtotal_capital_exigible", "Suma de compras a fin de mes y cuotas vencidas a amortizar a capital", "Float", "> 0", "S/ #,###.##"),
        ("total_a_pagar_fecha_pactada", "Importe total líquido exigible a pagar (Capital + Interés Comp. + Intereses Mora)", "Float", "total = capital + I_comp + I_mora", "S/ #,###.##"),
        ("fecha_vencimiento_pago", "Fecha límite pactada para el pago sin recargo de intereses moratorios", "Date", "Fecha válida del ciclo actual", "DD/MM/AAAA"),
        ("estado_liquidacion_ciclo", "Estado de cumplimiento de la liquidación generada en el corte", "Enum", "{Pendiente, Pagado_Puntual, Pagado_Con_Mora, En_Mora}", "Texto"),
        # Cronograma frances
        ("cronograma_cuotas_frances", "Tabla mensual de cuota, interés, amortización y saldo bajo método francés ordinario", "Tabla", "N filas (1 <= N <= PlazoMax)", "Cronograma N períodos"),
        ("cuota_mensual_fija (C)", "Cuota mensual constante a pagar cada 30 días según el método francés", "Float", "C > 0", "S/ #,###.##"),
        ("capital_financiado_capitalizado (P_cap)", "Deuda base para cuotas francesas tras capitalizar periodo de gracia total", "Float", "P_cap >= P_0", "S/ #,###.##"),
        ("total_intereses_financiados", "Suma total de intereses compensatorios en todas las cuotas del cronograma", "Float", "Total_I = (N * C) - P_cap >= 0", "S/ #,###.##"),
        ("monto_total_pagar_cuotas", "Importe acumulado total desembolsado por el cliente por la compra en cuotas", "Float", "Total_Pagar = N * C", "S/ #,###.##"),
        # Indicadores financieros
        ("TEM_calculada", "Tasa Efectiva Mensual comercial de 30 días aplicada en el cronograma", "Float", "0.0000000% <= TEM <= 10%", "#.#######%"),
        ("TED_calculada", "Tasa Efectiva Diaria aplicada en períodos de gracia e intereses diarios", "Float", "0.0000000% <= TED <= 0.5%", "#.#######%"),
        ("TCEA_operacion", "Tasa de Costo Efectivo Anual de la operación de crédito para el cliente", "Float", "TCEA >= 0", "#.####%"),
        ("TIR_operacion", "Tasa Interna de Retorno anualizada obtenida por el comercio prestamista", "Float", "TIR >= 0", "#.####%"),
        ("VAN_operacion", "Valor Actual Neto de los cobros futuros descontados al COK del comercio", "Float", "Número real (positivo/negativo)", "S/ #,###.##"),
        ("ganancia_bruta_intereses", "Ganancia monetaria neta del negocio por concepto de intereses cobrados", "Float", ">= 0", "S/ #,###.##"),
        # Estado de cuenta cliente
        ("saldo_deuda_total_acumulada", "Deuda global total pendiente consolidada que mantiene el cliente con la tienda", "Float", "0.00 <= Saldo <= limite_credito", "S/ #,###.##"),
        ("credito_disponible", "Margen remanente de línea de crédito disponible para compras futuras", "Float", "credito_disp = LC - saldo_deuda", "S/ #,###.##"),
        ("dias_mora_acumulados", "Días de retraso acumulados si el cliente no canceló a la fecha pactada", "Int", ">= 0", "#"),
        ("estado_crediticio_cliente", "Calificación del comportamiento de pago y estatus operativo del cliente", "Enum", "{Al_Día, Por_Vencer, En_Mora, Bloqueado}", "Texto")
    ]
    
    table_out = doc.add_table(rows=len(data_outputs)+1, cols=5)
    set_table_layout(table_out, out_col_widths)
    
    headers_out = ["VARIABLE", "DESCRIPCIÓN", "TIPO", "CONDICIÓN / RANGO", "FORMATO"]
    hdr_row_o = table_out.rows[0]
    hdr_trPr_o = hdr_row_o._tr.get_or_add_trPr()
    hdr_trPr_o.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for i, h_text in enumerate(headers_out):
        cell = hdr_row_o.cells[i]
        format_cell(cell, fill_hex="c8102e")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = WHITE
        
    for r_idx, row_data in enumerate(data_outputs):
        row = table_out.rows[r_idx+1]
        fill = "f2f2f2" if r_idx % 2 == 1 else "ffffff"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            format_cell(cell, fill_hex=fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True

    # -------------------------------------------------------------
    # 5. Datos Intermedios
    # -------------------------------------------------------------
    sub3 = doc.add_paragraph()
    sub3.paragraph_format.space_before = Pt(14)
    sub3.paragraph_format.space_after = Pt(4)
    run_sub3 = sub3.add_run('6.a.3. Datos Intermedios')
    run_sub3.font.size = Pt(13)
    run_sub3.font.bold = True
    run_sub3.font.color.rgb = UPC_RED
    
    p_desc_mid = doc.add_paragraph()
    p_desc_mid.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_desc_mid.paragraph_format.space_after = Pt(8)
    p_desc_mid.paragraph_format.line_spacing = 1.15
    p_desc_mid.add_run(
        "Los datos intermedios son variables auxiliares y de cómputo generadas por los motores matemáticos del software. "
        "Comprenden la cuantificación de días calendario y de gracia, la conversión estandarizada de tasas efectivas y nominales "
        "(con base de 360 días), la capitalización de la deuda en periodos de diferimiento inicial, el desdoblamiento periódico "
        "de cuotas francesas (interés y amortización), la liquidación de moras e imputaciones bajo estricto orden de prelación, "
        "así como los vectores de flujos netos requeridos para los algoritmos numéricos de TIR y VAN. A continuación se detallan:"
    )
    
    mid_col_widths = [1900, 3160, 900, 1900, 1500]  # total = 9360
    
    data_intermediates = [
        # Metricas temporales
        ("dias_transcurridos_compra_pago (d)", "Días calendario transcurridos desde fecha de compra hasta fecha pactada", "Int", "0 <= d <= 60 días", "#"),
        ("dias_gracia_total (d_gracia)", "Días transcurridos entre compra y primer vencimiento con gracia total", "Int", "0 <= d_gracia <= 30 días", "#"),
        ("dias_mora_retraso (d_mora)", "Días en exceso transcurridos entre fecha de pago pactada y abono real", "Int", "d_mora >= 0 días", "#"),
        # Tasas y factores
        ("factor_ted_comp (f_TED)", "Tasa Efectiva Diaria compensatoria: (1+TEA)^(1/360)-1 ó TNA/360", "Float", "f_TED >= 0 (mín. 7 decimales)", "#.#######%"),
        ("factor_tem_comp (f_TEM)", "Tasa Efectiva Mensual comercial: (1+TEA)^(30/360)-1", "Float", "f_TEM >= 0 (mín. 7 decimales)", "#.#######%"),
        ("factor_ted_mora (f_TED_mora)", "Tasa Efectiva Diaria moratoria: (1+TEA_mora)^(1/360)-1 ó TNA_mora/360", "Float", "f_TED_mora >= 0 (mín. 7 dec.)", "#.#######%"),
        ("factor_recuperacion_capital (FRC)", "Factor francés vencido ordinario: [TEM*(1+TEM)^N] / [(1+TEM)^N - 1]", "Float", "FRC > 0", "#.########"),
        # Gracia y capitalizacion
        ("interes_gracia_capitalizado (I_gracia)", "Interés compensatorio devengado en gracia total: P_0 * [(1+TED)^d_gracia - 1]", "Float", "I_gracia >= 0", "S/ #,###.##"),
        ("capital_capitalizado_base (P_cap)", "Saldo deudor capitalizado base para método francés: P_cap = P_0 + I_gracia", "Float", "P_cap >= P_0", "S/ #,###.##"),
        # Periodo k
        ("interes_periodo_k (I_k)", "Interés compensatorio de la cuota k: I_k = S_{k-1} * f_TEM", "Float", "0 <= I_k < C", "S/ #,###.##"),
        ("amortizacion_periodo_k (A_k)", "Amortización a capital de la cuota k: A_k = C - I_k", "Float", "0 < A_k <= C", "S/ #,###.##"),
        ("saldo_deudor_periodo_k (S_k)", "Capital vivo insoluto tras pagar cuota k: S_k = S_{k-1} - A_k (S_N = 0.00)", "Float", "0 <= S_k < S_{k-1}", "S/ #,###.##"),
        # Mora y prelacion
        ("base_imponible_mora (B_mora)", "Deuda líquida exigible vencida sobre la cual se calculan intereses moratorios", "Float", "B_mora >= 0", "S/ #,###.##"),
        ("monto_interes_mora_calculado (I_mora)", "Interés moratorio devengado: B_mora * [(1+TED_mora)^d_mora - 1]", "Float", "I_mora >= 0", "S/ #,###.##"),
        ("monto_imputado_mora", "Porción de pago abonada priorizada en orden 1 a cancelar intereses por mora", "Float", "0 <= imp <= I_mora", "S/ #,###.##"),
        ("monto_imputado_compensatorio", "Porción de pago abonada asignada en orden 2 a cancelar interés compensatorio", "Float", "0 <= imp <= I_comp", "S/ #,###.##"),
        ("monto_imputado_capital", "Porción residual de pago asignada en orden 3 a amortizar capital adeudado", "Float", "imp >= 0", "S/ #,###.##"),
        # Flujos y control
        ("vector_flujos_caja (CF)", "Vector de flujos netos periódicos [-P_0, C_1, C_2, ... C_N] para TIR y VAN", "Vector Float", "Longitud N+1 elementos", "[-P_0, C_1, ...]"),
        ("flag_compra_post_corte", "Bandera lógica: 1 si compra fue post corte (pasa a mes M+1); 0 si mes M", "Boolean / Int", "{0, 1}", "Binario"),
        ("validacion_limite_credito", "Control lógico que comprueba si (saldo_actual + nueva_compra) <= LC", "Boolean", "{Verdadero, Falso}", "Lógico"),
        ("validacion_pago_completo", "Control lógico que verifica monto_abonado == total_exigible (sin pagos parciales)", "Boolean", "{Verdadero, Falso}", "Lógico")
    ]
    
    table_mid = doc.add_table(rows=len(data_intermediates)+1, cols=5)
    set_table_layout(table_mid, mid_col_widths)
    
    headers_mid = ["VARIABLE", "DESCRIPCIÓN", "TIPO", "CONDICIÓN / TAMAÑO", "FORMATO"]
    hdr_row_m = table_mid.rows[0]
    hdr_trPr_m = hdr_row_m._tr.get_or_add_trPr()
    hdr_trPr_m.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for i, h_text in enumerate(headers_mid):
        cell = hdr_row_m.cells[i]
        format_cell(cell, fill_hex="c8102e")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = WHITE
        
    for r_idx, row_data in enumerate(data_intermediates):
        row = table_mid.rows[r_idx+1]
        fill = "f2f2f2" if r_idx % 2 == 1 else "ffffff"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            format_cell(cell, fill_hex=fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True
                
    doc.save(output_path)
    print(f"File successfully saved to {output_path}!")

if __name__ == "__main__":
    src = r"C:\Users\Frank\Documents\antigravity\happy-carson\stonks\Finanzas Trabajo Parcial.docx"
    dst = r"C:\Users\Frank\Documents\antigravity\happy-carson\stonks\Finanzas Trabajo Parcial.docx"
    create_seccion_6a(src, dst)
